#!/usr/bin/env python3
"""
rag_core.py - Shared RAG + LLM + input-source helpers for the TestCaseGENI web app.

Reuses the SAME Qdrant collection / embedding model / payload shape produced by
ingest_from_excel.py:
    payload = {key, summary, product, content, metadata{objective, precondition,
               steps, test_data, expected_result, priority, severity, status,
               component, folder, test_type, automation_status, ...}}

Inputs supported for generation:
    * Jira story ID      (JIRA_BASE_URL / JIRA_EMAIL / JIRA_TOKEN)
    * Figma URL          (FIGMA_TOKEN - optional; text + rendered frame image)
    * Figma screenshot   (image upload; used as vision input if the LLM supports it)
    * Design document    (PDF / DOCX / TXT / MD upload)

LLM providers (LLM_PROVIDER env):
    * openai  (default) - OPENAI_API_KEY, OPENAI_MODEL (default gpt-4o)
    * ollama            - OLLAMA_URL, OLLAMA_MODEL, OLLAMA_VISION=true|false
"""

from __future__ import annotations

import base64
import io
import json
import logging
import os
import re
import time
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

import requests
from dotenv import load_dotenv

# .env lives in the project root (one level above /webapp)
PROJECT_ROOT = Path(__file__).resolve().parent.parent
load_dotenv(PROJECT_ROOT / ".env")

import sys  # noqa: E402
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
import glossary  # noqa: E402  shared with ingest_from_excel.py (A/B = ABT = AB testing ...)

log = logging.getLogger("testcasegeni")

# --------------------------------------------------------------------------- #
# Configuration (all overridable via .env)
# --------------------------------------------------------------------------- #
QDRANT_URL = os.getenv("QDRANT_URL", "http://localhost:6333")
COLLECTION = os.getenv("QDRANT_COLLECTION", "jira_test_cases_all")
EMBED_MODEL = os.getenv("EMBEDDING_MODEL", "all-MiniLM-L6-v2")

LLM_PROVIDER = os.getenv("LLM_PROVIDER", "openai").lower()
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o")
OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434").rstrip("/")
OLLAMA_MODEL = (os.getenv("OLLAMA_MODEL") or "").strip() or "qwen-128k:latest"   # empty in .env -> default
OLLAMA_VISION = os.getenv("OLLAMA_VISION", "false").lower() == "true"
# Ollama's default context window is small and longer prompts are SILENTLY truncated. Generation prompts
# (story + design doc + references) need well over 8k tokens -> set it explicitly. Bigger = more RAM/VRAM.
OLLAMA_NUM_CTX = int(os.getenv("OLLAMA_NUM_CTX", "16384"))
OLLAMA_TIMEOUT = int(os.getenv("OLLAMA_TIMEOUT", "600"))   # local models are slower than the API

LLM_TIMEOUT = int(os.getenv("LLM_TIMEOUT", "180"))
LLM_MAX_RETRIES = int(os.getenv("LLM_MAX_RETRIES", "3"))

JIRA_BASE_URL = os.getenv("JIRA_BASE_URL")
JIRA_EMAIL = os.getenv("JIRA_EMAIL")
JIRA_TOKEN = os.getenv("JIRA_TOKEN")
FIGMA_TOKEN = os.getenv("FIGMA_TOKEN")

MAX_DOC_CHARS = int(os.getenv("MAX_DOC_CHARS", "60000"))  # docs are split into sections for analysis
import products  # noqa: E402  product list = config.json "products" (single source of truth)
KNOWN_PRODUCTS = products.product_ids()
STORY_ID_RE = re.compile(r"^[A-Z][A-Z0-9_]{1,19}-\d{1,9}$")

# --------------------------------------------------------------------------- #
# Lazy singletons (heavy objects load once, on first use)
# --------------------------------------------------------------------------- #
_embedder = None
_qdrant = None
_openai = None
_jira = None


def embedder():
    global _embedder
    if _embedder is None:
        from sentence_transformers import SentenceTransformer
        log.info("Loading embedding model %s", EMBED_MODEL)
        _embedder = SentenceTransformer(EMBED_MODEL)
    return _embedder


def qdrant():
    global _qdrant
    if _qdrant is None:
        from qdrant_client import QdrantClient
        _qdrant = QdrantClient(url=QDRANT_URL, api_key=os.getenv("QDRANT_API_KEY") or None,
                               prefer_grpc=False, timeout=10.0, check_compatibility=False)
    return _qdrant


def openai_client():
    global _openai
    if _openai is None:
        from openai import OpenAI
        if not os.getenv("OPENAI_API_KEY"):
            raise RuntimeError("OPENAI_API_KEY is not set in .env")
        _openai = OpenAI()
    return _openai


def jira_client():
    global _jira
    if _jira is None:
        if not (JIRA_BASE_URL and JIRA_EMAIL and JIRA_TOKEN):
            raise RuntimeError("Jira is not configured (JIRA_BASE_URL / JIRA_EMAIL / JIRA_TOKEN)")
        from jira import JIRA
        _jira = JIRA(server=JIRA_BASE_URL, basic_auth=(JIRA_EMAIL, JIRA_TOKEN), timeout=20)
    return _jira


# --------------------------------------------------------------------------- #
# Health / stats
# --------------------------------------------------------------------------- #
def health() -> dict:
    status = {"qdrant": False, "llm": False, "jira": bool(JIRA_BASE_URL and JIRA_TOKEN),
              "figma": bool(FIGMA_TOKEN), "provider": LLM_PROVIDER,
              "model": OPENAI_MODEL if LLM_PROVIDER == "openai" else OLLAMA_MODEL,
              "collection": COLLECTION, "points": None}
    try:
        info = qdrant().get_collection(COLLECTION)
        status["qdrant"] = True
        status["points"] = info.points_count
    except Exception as e:  # noqa: BLE001
        status["qdrant_error"] = str(e)[:200]
    if LLM_PROVIDER == "openai":
        status["llm"] = bool(os.getenv("OPENAI_API_KEY"))
    else:
        try:
            status["model"] = resolve_ollama_model(_ollama_installed())
            status["llm"] = True
        except LLMConfigError as e:
            status["llm_error"] = str(e)
        except requests.RequestException as e:
            status["llm_error"] = f"Ollama not reachable at {OLLAMA_URL}: {str(e)[:150]}"
    return status


def product_counts() -> dict:
    """Exact point count per product (small collection -> cheap)."""
    from qdrant_client.models import FieldCondition, Filter, MatchValue
    out = {}
    for p in KNOWN_PRODUCTS:
        try:
            c = qdrant().count(COLLECTION, count_filter=Filter(
                must=[FieldCondition(key="product", match=MatchValue(value=p))]), exact=True)
            out[p] = c.count
        except Exception:  # noqa: BLE001
            out[p] = None
    return out


# --------------------------------------------------------------------------- #
# Retrieval
# --------------------------------------------------------------------------- #
def _flatten(payload: dict, score: float) -> dict:
    md = payload.get("metadata", {}) or {}
    clean = lambda v: "" if v in (None, "nan", "None") else str(v)  # noqa: E731
    return {
        "key": clean(payload.get("key")),
        "title": clean(payload.get("summary")),
        "product": clean(payload.get("product")),
        "score": round(float(score), 4),
        "objective": clean(md.get("objective")),
        "precondition": clean(md.get("precondition")),
        "steps": clean(md.get("steps")),
        "test_data": clean(md.get("test_data")),
        "expected_result": clean(md.get("expected_result")),
        "priority": clean(md.get("priority")),
        "severity": clean(md.get("severity")),
        "component": clean(md.get("component")),
        "folder": clean(md.get("folder")),
        "test_type": clean(md.get("test_type")),
        "automation_status": clean(md.get("automation_status")),
        "features": list(payload.get("features") or []),
        "labels": clean(md.get("labels")),
    }


def feature_counts() -> list[dict]:
    """Glossary features + how many points are tagged with each (needs payload['features'])."""
    from qdrant_client.models import FieldCondition, Filter, MatchValue
    out = []
    for f in glossary.list_features():
        try:
            n = qdrant().count(COLLECTION, count_filter=Filter(
                must=[FieldCondition(key="features", match=MatchValue(value=f["id"]))]), exact=True).count
        except Exception:  # noqa: BLE001
            n = None
        out.append({**f, "count": n})
    return out


RRF_K = 60  # standard Reciprocal Rank Fusion constant


def retrieve(query: str, top_k: int = 5, product: str | None = None,
             min_score: float = 0.0, feature: str | None = None, expand: bool = True) -> list[dict]:
    """Semantic search over ingested test cases.
    product/feature = optional hard filters. expand=True appends glossary aliases to the query
    (e.g. "ABT sampling" also carries "A/B Testing, AB testing, ...") so every naming variant is found."""
    from qdrant_client.models import FieldCondition, Filter, MatchValue
    top_k = max(1, min(int(top_k), 50))
    q = query[:8000]
    intents = glossary.detect_intents(q) if expand else []
    if expand:
        q = glossary.rewrite_for_search(q)[0]   # QA intents, e.g. "negative test" -> invalid/error wording
        q = glossary.expand_query(q, product)[0]  # feature synonyms, e.g. ABT -> A/B Testing ...
    vec = embedder().encode(q).tolist()
    must = []
    if product and product.upper() in KNOWN_PRODUCTS:
        must.append(FieldCondition(key="product", match=MatchValue(value=product.upper())))
    valid_features = {f["id"] for f in glossary.list_features()}
    if feature and feature.upper() in valid_features:
        must.append(FieldCondition(key="features", match=MatchValue(value=feature.upper())))
    flt = Filter(must=must) if must else None
    if intents:
        # Two searches fused with Reciprocal Rank Fusion (RRF):
        #   topic  = question with the intent phrase removed ("... for Event Pattern")
        #   intent = question with the phrase rewritten ("... invalid input error ... Event Pattern")
        # Cases that rank well in BOTH (on-topic AND negative-like) come first.
        topic_q = query[:8000]
        for iid in intents:
            topic_q = glossary.load_intents()[iid]["regex"].sub(" ", topic_q)
        topic_vec = embedder().encode(glossary.expand_query(topic_q, product)[0]).tolist()
        pool = max(top_k * 4, 20)
        searches = [(vec, flt), (topic_vec, flt)]
        # 3rd list: on-topic cases carrying the intent's Zephyr label (e.g. Labels = "Negative")
        for iid in intents:
            tag = glossary.load_intents()[iid].get("label_tag")
            if tag:
                searches.append((topic_vec, Filter(must=must + [
                    FieldCondition(key="tags", match=MatchValue(value=tag))])))
        lists = [qdrant().query_points(collection_name=COLLECTION, query=v, limit=pool, query_filter=f,
                                       with_payload=True, score_threshold=min_score or None).points
                 for v, f in searches]
        fused, best = {}, {}
        for pts in lists:
            for rank, p in enumerate(pts):
                fused[p.id] = fused.get(p.id, 0.0) + 1.0 / (RRF_K + rank + 1)
                if p.id not in best or p.score > best[p.id].score:
                    best[p.id] = p
        order = sorted(fused, key=lambda i: -fused[i])[:top_k]
        return [_flatten(best[i].payload or {}, best[i].score) for i in order]

    res = qdrant().query_points(collection_name=COLLECTION, query=vec, limit=top_k,
                                query_filter=flt, with_payload=True,
                                score_threshold=min_score or None)
    return [_flatten(p.payload or {}, p.score) for p in res.points]


def format_context(cases: list[dict], max_chars_each: int = 1500) -> str:
    if not cases:
        return "(no matching test cases found in the knowledge base)"
    blocks = []
    for i, c in enumerate(cases, 1):
        b = (f"[{i}] {c['key']} | {c['title']} | product={c['product']} | component={c['component']} "
             f"| type={c['test_type']} | labels={c.get('labels') or '-'} | priority={c['priority']} "
             f"| similarity={c['score']}\n"
             f"Objective: {c['objective']}\nPrecondition: {c['precondition']}\n"
             f"Steps: {c['steps']}\nTest data: {c['test_data']}\nExpected: {c['expected_result']}")
        blocks.append(b[:max_chars_each])
    return "\n\n".join(blocks)


# --------------------------------------------------------------------------- #
# LLM abstraction
# --------------------------------------------------------------------------- #
class LLMConfigError(RuntimeError):
    """Non-transient LLM problem (wrong model name, bad request) - retrying cannot help."""


_resolved_model: str | None = None


def _ollama_installed() -> list[str]:
    r = requests.get(f"{OLLAMA_URL}/api/tags", timeout=5)
    r.raise_for_status()
    return [m.get("name", "") for m in (r.json().get("models") or [])]


def resolve_ollama_model(names: list[str] | None = None) -> str:
    """Exact name if installed; 'deepseek-r1' -> the single installed 'deepseek-r1:<tag>'; else a clear error."""
    global _resolved_model
    if _resolved_model and names is None:
        return _resolved_model
    names = _ollama_installed() if names is None else names
    want = OLLAMA_MODEL if ":" in OLLAMA_MODEL else OLLAMA_MODEL + ":latest"
    if want in names:
        found = want
    else:
        base = OLLAMA_MODEL.split(":")[0]
        cands = [n for n in names if n.split(":")[0] == base]
        if len(cands) == 1:
            found = cands[0]
            log.info("OLLAMA_MODEL=%s not installed as '%s' - using '%s'", OLLAMA_MODEL, want, found)
        else:
            hint = f"matching: {', '.join(cands)}" if cands else f"installed: {', '.join(names) or 'none'}"
            raise LLMConfigError(f"Ollama model '{OLLAMA_MODEL}' not found ({hint}). "
                                 f"Set OLLAMA_MODEL in .env to a name from 'ollama list'.")
    _resolved_model = found
    return found


_THINK = re.compile(r"<think>.*?</think>", re.DOTALL | re.IGNORECASE)


def strip_reasoning(text: str) -> str:
    """Remove reasoning blocks of 'thinking' models (DeepSeek-R1, QwQ ...): <think>...</think>."""
    text = _THINK.sub("", text or "")
    if "</think>" in text.lower():                   # opening tag missing / truncated start
        text = re.split(r"</think>", text, flags=re.IGNORECASE)[-1]
    return text.strip()


def llm_chat(messages: list[dict], json_mode: bool = False, images: list[str] | None = None,
             temperature: float = 0.3, max_tokens: int = 4000) -> str:
    """
    messages: [{"role": "system"|"user"|"assistant", "content": str}]
    images:   list of data-URLs or https URLs, attached to the LAST user message.
    Retries with exponential backoff on transient errors.
    """
    images = images or []
    last_err: Exception | None = None
    for attempt in range(1, LLM_MAX_RETRIES + 1):
        try:
            if LLM_PROVIDER == "openai":
                return _openai_chat(messages, json_mode, images, temperature, max_tokens)
            return _ollama_chat(messages, json_mode, images, temperature, max_tokens)
        except LLMConfigError:
            raise                                   # wrong model / bad request: fail fast, no retries
        except Exception as e:  # noqa: BLE001
            last_err = e
            log.warning("LLM attempt %d/%d failed: %s", attempt, LLM_MAX_RETRIES, e)
            if attempt < LLM_MAX_RETRIES:
                time.sleep(2 ** attempt)
    raise RuntimeError(f"LLM call failed after {LLM_MAX_RETRIES} attempts: {last_err}")


def _openai_chat(messages, json_mode, images, temperature, max_tokens) -> str:
    msgs = [dict(m) for m in messages]
    if images:
        last = msgs[-1]
        parts = [{"type": "text", "text": last["content"]}]
        parts += [{"type": "image_url", "image_url": {"url": u}} for u in images]
        last["content"] = parts
    kwargs: dict[str, Any] = dict(model=OPENAI_MODEL, messages=msgs, temperature=temperature,
                                  max_tokens=max_tokens, timeout=LLM_TIMEOUT)
    if json_mode:
        kwargs["response_format"] = {"type": "json_object"}
    r = openai_client().chat.completions.create(**kwargs)
    return (r.choices[0].message.content or "").strip()


def _ollama_chat(messages, json_mode, images, temperature, max_tokens=4000) -> str:
    msgs = [dict(m) for m in messages]
    if images and OLLAMA_VISION:
        b64 = []
        for u in images:
            if u.startswith("data:"):
                b64.append(u.split(",", 1)[1])
            else:
                b64.append(base64.b64encode(requests.get(u, timeout=20).content).decode())
        msgs[-1]["images"] = b64
    body = {"model": resolve_ollama_model(), "messages": msgs, "stream": False,
            "options": {"temperature": temperature, "num_ctx": OLLAMA_NUM_CTX, "num_predict": max_tokens}}
    if json_mode:
        body["format"] = "json"
    r = requests.post(f"{OLLAMA_URL}/api/chat", json=body, timeout=OLLAMA_TIMEOUT)
    if not r.ok:
        try:
            detail = r.json().get("error", r.text)
        except ValueError:
            detail = r.text
        msg = f"Ollama HTTP {r.status_code}: {str(detail)[:300]}"
        if r.status_code in (400, 404):
            global _resolved_model
            _resolved_model = None                  # re-resolve next time (model may have been pulled)
            raise LLMConfigError(msg)
        raise RuntimeError(msg)
    data = r.json()
    # Warn when the prompt filled the whole window (= likely truncated by Ollama)
    if data.get("prompt_eval_count", 0) >= OLLAMA_NUM_CTX - 16:
        log.warning("Ollama prompt used %s tokens = full num_ctx %s -> input probably truncated; "
                    "raise OLLAMA_NUM_CTX", data.get("prompt_eval_count"), OLLAMA_NUM_CTX)
    return strip_reasoning(data.get("message", {}).get("content", ""))


def _loads_lenient(t: str) -> Any:
    """json.loads that tolerates what local models often emit: raw newlines/tabs inside strings
    (strict=False) and trailing commas before } or ]."""
    try:
        return json.loads(t, strict=False)
    except json.JSONDecodeError:
        return json.loads(re.sub(r",\s*([}\]])", r"\1", t), strict=False)


def parse_json(text: str) -> Any:
    """Robust JSON extraction (reasoning blocks, fences, leading prose, raw newlines, trailing commas)."""
    t = strip_reasoning(text or "")
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", t.strip(), flags=re.MULTILINE)
    try:
        return _loads_lenient(t)
    except json.JSONDecodeError:
        pass
    for open_c, close_c in (("{", "}"), ("[", "]")):
        a, b = t.find(open_c), t.rfind(close_c)
        if a != -1 and b > a:
            try:
                return _loads_lenient(t[a:b + 1])
            except json.JSONDecodeError:
                continue
    raise ValueError("LLM did not return valid JSON")


def salvage_answer(text: str) -> str:
    """Last resort for chat: pull the answer out of broken JSON, or accept plain Markdown."""
    t = strip_reasoning(text or "").strip()
    m = re.search(r'"answer"\s*:\s*"(.*)', t, re.DOTALL)
    if m:                                            # broken JSON: take the answer string up to its end
        body = re.sub(r'"\s*,?\s*("in_scope"\s*:\s*\w+\s*)?}?\s*$', "", m.group(1).strip())
        return body.replace('\\n', "\n").replace('\\"', '"').replace("\\t", " ").strip()
    if t and not t.lstrip().startswith(("{", "[")):  # plain Markdown answer, no JSON at all
        return t
    return ""


# --------------------------------------------------------------------------- #
# Input sources
# --------------------------------------------------------------------------- #
def _adf_to_text(node: Any) -> str:
    """Recursively flatten Atlassian Document Format (handles lists, tables, headings)."""
    if isinstance(node, str):
        return node
    if isinstance(node, list):
        return "".join(_adf_to_text(n) for n in node)
    if not isinstance(node, dict):
        return ""
    t = node.get("type")
    if t == "text":
        return node.get("text", "")
    inner = _adf_to_text(node.get("content", []))
    if t in ("paragraph", "heading", "listItem", "tableRow", "codeBlock", "blockquote"):
        return inner + "\n"
    if t == "tableCell" or t == "tableHeader":
        return inner.strip() + " | "
    if t == "hardBreak":
        return "\n"
    return inner


def fetch_jira_story(story_id: str) -> dict:
    story_id = story_id.strip().upper()
    if not STORY_ID_RE.match(story_id):
        raise ValueError(f"Invalid story ID format: {story_id!r} (expected e.g. HMA-401957)")
    issue = jira_client().issue(story_id)
    f = issue.fields
    desc = f.description
    desc_text = _adf_to_text(desc) if isinstance(desc, (dict, list)) else (desc or "")
    return {
        "key": story_id,
        "summary": f.summary or "",
        "type": getattr(f.issuetype, "name", ""),
        "status": getattr(f.status, "name", ""),
        "components": [c.name for c in (getattr(f, "components", None) or [])],
        "labels": list(getattr(f, "labels", None) or []),
        "description": desc_text[:MAX_DOC_CHARS],
        "url": f"{JIRA_BASE_URL.rstrip('/')}/browse/{story_id}",
    }


def parse_figma_url(url: str) -> tuple[str, str | None]:
    """Return (file_key, node_id). Only figma.com hosts accepted (prevents SSRF)."""
    u = urlparse(url.strip())
    host = (u.hostname or "").lower()
    if u.scheme != "https" or not (host == "figma.com" or host.endswith(".figma.com")):
        raise ValueError("Figma link must be an https://www.figma.com/... URL")
    m = re.match(r"^/(?:file|design|proto|board)/([A-Za-z0-9]+)", u.path)
    if not m:
        raise ValueError("Could not find a Figma file key in the URL")
    node = parse_qs(u.query).get("node-id", [None])[0]
    return m.group(1), (node.replace("-", ":") if node else None)


def fetch_figma(url: str) -> dict:
    """Pull frame names + all text layers, plus a rendered PNG URL of the node."""
    if not FIGMA_TOKEN:
        raise RuntimeError("FIGMA_TOKEN not set - upload a screenshot of the Figma screen instead")
    key, node = parse_figma_url(url)
    h = {"X-Figma-Token": FIGMA_TOKEN}
    if node:
        r = requests.get(f"https://api.figma.com/v1/files/{key}/nodes", params={"ids": node, "depth": 50},
                         headers=h, timeout=30)
        r.raise_for_status()
        docs = [v["document"] for v in r.json().get("nodes", {}).values() if v]
    else:
        r = requests.get(f"https://api.figma.com/v1/files/{key}", params={"depth": 4}, headers=h, timeout=30)
        r.raise_for_status()
        docs = [r.json().get("document", {})]

    texts, frames = [], []

    def walk(n, depth=0):
        if n.get("type") in ("FRAME", "COMPONENT", "INSTANCE", "SECTION") and depth <= 3:
            frames.append(n.get("name", ""))
        if n.get("type") == "TEXT" and n.get("characters"):
            texts.append(n["characters"].strip())
        for c in n.get("children", []) or []:
            walk(c, depth + 1)

    for d in docs:
        walk(d)

    image_url = None
    if node:
        try:
            ir = requests.get(f"https://api.figma.com/v1/images/{key}",
                              params={"ids": node, "format": "png", "scale": 1}, headers=h, timeout=30)
            if ir.ok:
                image_url = next(iter((ir.json().get("images") or {}).values()), None)
        except requests.RequestException:
            pass

    seen, uniq = set(), []
    for t in texts:
        if t not in seen:
            seen.add(t)
            uniq.append(t)
    body = "Frames/screens: " + ", ".join(f for f in frames if f)[:2000] + \
           "\nUI text / labels:\n- " + "\n- ".join(uniq)
    return {"file_key": key, "node_id": node, "text": body[:MAX_DOC_CHARS], "image_url": image_url}


def extract_document_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
    elif ext == ".docx":
        import docx  # python-docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext in (".txt", ".md", ".csv", ".json"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported design-doc type {ext} (use PDF, DOCX, TXT, MD)")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:MAX_DOC_CHARS]


def image_to_data_url(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    mime = {"jpg": "jpeg", "jpeg": "jpeg", "png": "png", "webp": "webp", "gif": "gif"}.get(ext)
    if not mime:
        raise ValueError("Screenshot must be PNG, JPG, WEBP or GIF")
    return f"data:image/{mime};base64,{base64.b64encode(data).decode()}"
