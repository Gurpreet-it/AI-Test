#!/usr/bin/env python3
"""
server.py - FastAPI backend + static UI for TestCaseGENI.

Run (from project root, venv active):
    pip install -r webapp/requirements.txt
    python3 webapp/server.py            # -> http://127.0.0.1:8080

Endpoints
    GET  /                 UI (webapp/static/index.html)
    GET  /api/health       Qdrant / LLM / Jira / Figma status
    GET  /api/products     test-case count per product in Qdrant
    POST /api/search       raw top-K retrieval (JSON)
    POST /api/chat         RAG chat (JSON)
    POST /api/generate     test-case generation (multipart form)
    POST /api/export       generated cases -> .xlsx (JSON)
"""

from __future__ import annotations

import io
import logging
import os
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

sys.path.insert(0, str(Path(__file__).resolve().parent))

from fastapi import FastAPI, File, Form, HTTPException, UploadFile  # noqa: E402
from fastapi.responses import FileResponse, StreamingResponse  # noqa: E402
from fastapi.staticfiles import StaticFiles  # noqa: E402
from starlette.concurrency import run_in_threadpool  # noqa: E402
from pydantic import BaseModel, Field  # noqa: E402

import rag_core as rc  # noqa: E402
import testcase_agent as agent  # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("testcasegeni.server")

STATIC = Path(__file__).resolve().parent / "static"
OUTPUT_DIR = rc.PROJECT_ROOT / "generated_test_cases"
MAX_UPLOAD = int(os.getenv("MAX_UPLOAD_MB", "15")) * 1024 * 1024

app = FastAPI(title="TestCaseGENI", version="1.0")
app.mount("/static", StaticFiles(directory=STATIC), name="static")


# ----------------------------- models -------------------------------------- #
class ChatTurn(BaseModel):
    role: str
    content: str


class ChatRequest(BaseModel):
    message: str = Field(..., min_length=1, max_length=4000)
    history: list[ChatTurn] = []
    product: Optional[str] = None
    feature: Optional[str] = None
    top_k: int = Field(5, ge=1, le=20)


class SearchRequest(BaseModel):
    query: str = Field(..., min_length=1, max_length=4000)
    product: Optional[str] = None
    feature: Optional[str] = None
    top_k: int = Field(5, ge=1, le=50)


class ExportRequest(BaseModel):
    story_id: Optional[str] = None
    test_cases: list[dict]
    references: list[dict] = []
    requirements: list[dict] = []
    open_questions: list[str] = []


# ----------------------------- helpers ------------------------------------- #
async def _read_upload(f: Optional[UploadFile]) -> Optional[tuple[str, bytes]]:
    if not f or not f.filename:
        return None
    data = await f.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, f"{f.filename} exceeds {MAX_UPLOAD // (1024 * 1024)} MB")
    return Path(f.filename).name, data  # strip any client-side path


def _fail(e: Exception, what: str):
    log.exception("%s failed", what)
    code = 400 if isinstance(e, ValueError) else 502
    raise HTTPException(code, f"{what} failed: {e}")


# ----------------------------- routes -------------------------------------- #
@app.get("/")
def index():
    return FileResponse(STATIC / "index.html")


@app.get("/api/health")
def api_health():
    return rc.health()


@app.get("/api/products")
def api_products():
    return {"products": rc.product_counts()}


@app.get("/api/features")
def api_features():
    """Glossary features (A/B Testing ...) with tagged-point counts."""
    try:
        return {"features": rc.feature_counts()}
    except Exception as e:  # noqa: BLE001
        _fail(e, "Features")


@app.post("/api/search")
def api_search(req: SearchRequest):
    try:
        return {"results": rc.retrieve(req.query, req.top_k, req.product, feature=req.feature)}
    except Exception as e:  # noqa: BLE001
        _fail(e, "Search")


@app.post("/api/chat")
def api_chat(req: ChatRequest):
    try:
        return agent.chat_answer(req.message, [t.model_dump() for t in req.history], req.product, req.top_k,
                                 feature=req.feature)
    except Exception as e:  # noqa: BLE001
        _fail(e, "Chat")


@app.post("/api/generate")
async def api_generate(
    story_id: str = Form(""),
    figma_url: str = Form(""),
    notes: str = Form(""),
    product: str = Form(""),
    feature: str = Form(""),
    top_k: int = Form(5),
    num_cases: int = Form(0),   # 0 = auto / full coverage
    design_doc: Optional[UploadFile] = File(None),
    screenshot: Optional[UploadFile] = File(None),
):
    doc = await _read_upload(design_doc)
    shot = await _read_upload(screenshot)
    try:
        # Generation takes minutes and is blocking -> run in a worker thread so the event loop
        # (chat, health, other users) keeps serving requests meanwhile.
        result = await run_in_threadpool(
            agent.generate_testcases,
            story_id=story_id.strip() or None, figma_url=figma_url.strip() or None,
            design_doc=doc, screenshot=shot, notes=notes, product=product or None,
            top_k=max(1, min(top_k, 20)), num_cases=num_cases or None,
            feature=feature or None)
    except Exception as e:  # noqa: BLE001
        _fail(e, "Generation")
    # Audit trail: keep a JSON copy of every run
    try:
        import json
        OUTPUT_DIR.mkdir(exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        name = f"{agent.safe_story_id(story_id)}_{stamp}.json"
        (OUTPUT_DIR / name).write_text(json.dumps(result, indent=2, ensure_ascii=False))
        result["saved_as"] = f"generated_test_cases/{name}"
    except OSError as e:
        log.warning("Could not save run: %s", e)
    return result


@app.post("/api/export")
def api_export(req: ExportRequest):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    head_fill = PatternFill("solid", start_color="17707F")
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    alt_fill = PatternFill("solid", start_color="F7F7FC")
    thin = Side(style="thin", color="DCE6F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Generated Test Cases"
    cols = ["TC ID", "Story ID", "Title", "Objective", "Preconditions", "Step #", "Step", "Test Data",
            "Step Expected Result", "Overall Expected Result", "Priority", "Test Type", "Component",
            "Requirement IDs", "Traceability", "Reference Keys", "Has TBD"]
    ws.append(cols)
    row_i = 0
    for tc in req.test_cases:
        steps = tc.get("steps") or [{}]
        for n, s in enumerate(steps, 1):
            first = n == 1
            ws.append([tc.get("tc_id", ""), req.story_id or "", tc.get("title", "") if first else "",
                       tc.get("objective", "") if first else "", tc.get("preconditions", "") if first else "",
                       n, s.get("step", ""), s.get("test_data", ""), s.get("expected", ""),
                       tc.get("expected_result", "") if first else "", tc.get("priority", "") if first else "",
                       tc.get("test_type", "") if first else "", tc.get("component", "") if first else "",
                       ", ".join(tc.get("requirement_ids", []) or []) if first else "",
                       tc.get("traceability", "") if first else "",
                       ", ".join(tc.get("reference_keys", []) or []) if first else "",
                       ("YES" if tc.get("has_tbd") else "") if first else ""])
            if row_i % 2:
                for c in ws[ws.max_row]:
                    c.fill = alt_fill
        row_i += 1

    ws2 = wb.create_sheet("RAG References")
    ws2.append(["Key", "Title", "Product", "Component", "Similarity", "Objective"])
    for r in req.references:
        ws2.append([r.get("key"), r.get("title"), r.get("product"), r.get("component"),
                    r.get("score"), r.get("objective")])

    ws3 = wb.create_sheet("Requirement Coverage")
    ws3.append(["Req ID", "Source", "Category", "Requirement", "Evidence (from source)", "Priority",
                "Covered", "Test Cases"])
    for r in req.requirements:
        ws3.append([r.get("id"), r.get("source"), r.get("category"), r.get("text"), r.get("evidence"),
                    r.get("priority"), "YES" if r.get("covered") else "NO", ", ".join(r.get("test_cases", []))])
    ws4 = wb.create_sheet("Open Questions")
    ws4.append(["#", "Question / missing information (resolve before execution)"])
    for i, q in enumerate(req.open_questions, 1):
        ws4.append([i, q])

    widths = [18, 14, 40, 40, 32, 7, 45, 22, 40, 36, 10, 14, 18, 16, 30, 22, 9]
    for sheet, w in ((ws, widths), (ws2, [16, 45, 12, 20, 11, 60]),
                     (ws3, [10, 26, 14, 50, 50, 10, 9, 40]), (ws4, [5, 110])):
        for i, cell in enumerate(sheet[1]):
            cell.fill, cell.font = head_fill, head_font
            sheet.column_dimensions[cell.column_letter].width = w[i]
        for row in sheet.iter_rows(min_row=1):
            for c in row:
                c.border, c.alignment = border, wrap
                if c.row > 1:
                    c.font = Font(name="Arial", size=10)
        sheet.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"test_cases_{agent.safe_story_id(req.story_id)}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


if __name__ == "__main__":
    import uvicorn
    host = os.getenv("APP_HOST", "127.0.0.1")  # local-only by default
    port = int(os.getenv("APP_PORT", "8080"))
    print(f"\n  TestCaseGENI UI  ->  http://{host}:{port}\n")
    uvicorn.run(app, host=host, port=port)
